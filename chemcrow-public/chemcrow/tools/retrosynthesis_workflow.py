"""Agent tools for an explicit, evaluated retrosynthesis workflow.

The tools in this module expose the planner's scoring and AND-OR search stages
without asking the language model to perform graph traversal itself.
"""

from __future__ import annotations

import csv
import heapq
import json
import math
from pathlib import Path
from typing import Any, Dict, List, Mapping, Optional, Sequence, Tuple

from langchain.pydantic_v1 import PrivateAttr
from langchain.tools import BaseTool

from chemcrow.planner import (
    EnzRetroClient,
    EnzymeReactionDatabase,
    HybridPlannerEngine,
    PlannerConfig,
)
from chemcrow.planner.chem_utils import (
    classify_precursors,
    extract_ec_numbers,
)
from chemcrow.planner.models import MoleculeNode, ReactionCandidate
from chemcrow.planner.scoring import score_and_rank_candidates
from chemcrow.tools.enzyme_planner import _parse_planner_input

__all__ = [
    "EnzymeSingleStepEvaluator",
    "EnzymeRetrosynthesisSearch",
    "EnzymeMultiStepEvaluator",
]

_DEFAULT_DB_PATH = str(
    Path(__file__).resolve().parents[1] / "data" / "enzyme_retro.db"
)


def _error(error_type: str, message: str) -> str:
    return json.dumps(
        {"ok": False, "error": {"type": error_type, "message": message}},
        ensure_ascii=False,
    )


def _json_object(raw_input: str) -> Dict[str, Any]:
    text = str(raw_input or "").strip()
    if not text:
        raise ValueError("Input cannot be empty.")
    try:
        payload = json.loads(text)
    except json.JSONDecodeError as exc:
        raise ValueError(f"Invalid JSON input: {exc.msg}.") from exc
    if not isinstance(payload, dict):
        raise ValueError("JSON input must be an object.")
    return payload


def _as_float(value: Any) -> Optional[float]:
    if value is None:
        return None
    if hasattr(value, "item"):
        value = value.item()
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _clamp(value: float) -> float:
    return max(0.0, min(1.0, float(value)))


def _string_list(value: Any) -> List[str]:
    if value is None:
        return []
    if isinstance(value, str):
        return [item for item in value.split(".") if item]
    if isinstance(value, (list, tuple)):
        result = []
        for item in value:
            if isinstance(item, Mapping):
                query = (
                    item.get("canonical_smiles")
                    or item.get("smiles")
                    or item.get("id")
                    or item.get("name")
                )
            else:
                query = item
            if query is not None and str(query).strip():
                result.append(str(query).strip())
        return result
    raise ValueError("A precursor list must be a list or dot-separated SMILES string.")


def _query_from_molecule(value: Any) -> str:
    if isinstance(value, Mapping):
        value = (
            value.get("canonical_smiles")
            or value.get("smiles")
            or value.get("id")
            or value.get("name")
        )
    text = str(value or "").strip()
    if not text:
        raise ValueError("A molecule must contain a SMILES, database ID, or exact name.")
    return text


def _route_signature(route: Mapping[str, Any]) -> Tuple[str, ...]:
    values = []
    for index, step in enumerate(route.get("steps", [])):
        candidate_id = step.get("candidate_id")
        if not candidate_id:
            database_ids = step.get("database_reaction_ids") or []
            candidate_id = "|".join(str(value) for value in database_ids)
        if not candidate_id:
            target = step.get("target") or {}
            precursors = step.get("precursors") or []
            candidate_id = json.dumps(
                [target, precursors, index], sort_keys=True, ensure_ascii=False
            )
        values.append(str(candidate_id))
    return tuple(sorted(values))


def _jaccard(left: Sequence[str], right: Sequence[str]) -> float:
    a, b = set(left), set(right)
    if not a and not b:
        return 1.0
    if not a or not b:
        return 0.0
    return len(a & b) / len(a | b)


_LIBRARY_COLUMNS = (
    "canonical_smiles",
    "smiles",
    "substrate_smiles",
    "building_block",
    "molecule",
    "compound",
)


def _library_record_value(value: Any, column: Optional[str] = None) -> Optional[str]:
    if isinstance(value, Mapping):
        keys = [column] if column else []
        keys.extend(_LIBRARY_COLUMNS)
        for key in keys:
            if key and value.get(key) is not None and str(value.get(key)).strip():
                return str(value.get(key)).strip()
        return None
    text = str(value or "").strip()
    return text or None


def _load_building_block_library(value: Any) -> Tuple[List[str], Dict[str, Any]]:
    """Load a local stock library without sending the whole library through the LLM."""
    if isinstance(value, (list, tuple)):
        items = [item for item in (_library_record_value(item) for item in value) if item]
        return items, {"type": "inline", "entry_count": len(items)}
    column = None
    sheet = None
    if isinstance(value, Mapping):
        path_value = value.get("path") or value.get("file")
        column = str(value.get("column") or "").strip() or None
        sheet = str(value.get("sheet") or "").strip() or None
    else:
        path_value = value
    path = Path(str(path_value or "")).expanduser().resolve()
    if not path.is_file():
        raise ValueError(f"Building-block library does not exist: {path}")
    suffix = path.suffix.casefold()
    items: List[str] = []

    if suffix in {".txt", ".smi", ".smiles"}:
        for line in path.read_text(encoding="utf-8").splitlines():
            stripped = line.strip()
            if stripped and not stripped.startswith("#"):
                # Plain TXT may contain exact local names with spaces. SMI/SMILES
                # files conventionally store SMILES in the first whitespace column.
                items.append(stripped if suffix == ".txt" else stripped.split()[0])
    elif suffix in {".csv", ".tsv"}:
        delimiter = "\t" if suffix == ".tsv" else ","
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle, delimiter=delimiter)
            fieldnames = list(reader.fieldnames or ())
            selected = column or next(
                (name for name in _LIBRARY_COLUMNS if name in fieldnames),
                fieldnames[0] if fieldnames else None,
            )
            if selected is None or selected not in fieldnames:
                raise ValueError(
                    f"No usable SMILES column in library {path}; columns={fieldnames}"
                )
            for row in reader:
                item = _library_record_value(row, selected)
                if item:
                    items.append(item)
            column = selected
    elif suffix in {".json", ".jsonl"}:
        if suffix == ".jsonl":
            records = [
                json.loads(line)
                for line in path.read_text(encoding="utf-8").splitlines()
                if line.strip()
            ]
        else:
            records = json.loads(path.read_text(encoding="utf-8"))
            if isinstance(records, Mapping):
                records = (
                    records.get("building_blocks")
                    or records.get("substrates")
                    or records.get("molecules")
                    or records.get("smiles")
                    or []
                )
        if not isinstance(records, list):
            raise ValueError("JSON building-block library must contain a list.")
        for record in records:
            item = _library_record_value(record, column)
            if item:
                items.append(item)
    elif suffix in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("openpyxl is required to read an XLSX library.") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        worksheet = workbook[sheet] if sheet else workbook.active
        rows = worksheet.iter_rows(values_only=True)
        headers = [str(item or "").strip() for item in next(rows, ())]
        selected = column or next(
            (name for name in _LIBRARY_COLUMNS if name in headers),
            headers[0] if headers else None,
        )
        if selected is None or selected not in headers:
            raise ValueError(
                f"No usable SMILES column in library {path}; columns={headers}"
            )
        index = headers.index(selected)
        for row in rows:
            if index < len(row) and row[index] is not None and str(row[index]).strip():
                items.append(str(row[index]).strip())
        column = selected
        sheet = worksheet.title
        workbook.close()
    else:
        raise ValueError(
            "Unsupported building-block library format; use TXT/SMI, CSV/TSV, "
            "JSON/JSONL, or XLSX."
        )

    items = list(dict.fromkeys(items))
    if not items:
        raise ValueError(f"Building-block library is empty: {path}")
    return items, {
        "type": "file",
        "path": str(path),
        "format": suffix.lstrip("."),
        "column": column,
        "sheet": sheet,
        "entry_count": len(items),
    }


def _parse_search_input(
    raw_input: str,
) -> Tuple[List[str], List[str], Dict[str, Any], Dict[str, Any]]:
    text = str(raw_input or "").strip()
    if not text.startswith("{"):
        target, blocks, options = _parse_planner_input(raw_input)
        return [target], blocks, options, {"library_sources": []}
    payload = _json_object(text)
    target_values = payload.get("targets")
    if target_values is None:
        target_values = payload.get("target") or payload.get("target_smiles")
    if isinstance(target_values, (str, Mapping)):
        target_values = [target_values]
    if not isinstance(target_values, (list, tuple)) or not target_values:
        raise ValueError("target/targets must contain at least one molecule.")
    targets = list(dict.fromkeys(_query_from_molecule(item) for item in target_values))

    block_values = payload.get("building_blocks", payload.get("substrates", []))
    if isinstance(block_values, str):
        delimiter = "|" if "|" in block_values else ","
        blocks = [item.strip() for item in block_values.split(delimiter) if item.strip()]
    elif isinstance(block_values, (list, tuple)):
        blocks = [_query_from_molecule(item) for item in block_values]
    elif block_values is None:
        blocks = []
    else:
        raise ValueError("building_blocks/substrates must be a list or string.")

    library_sources = []
    library_value = payload.get("building_block_library", payload.get("substrate_library"))
    if library_value is not None:
        library_values, library_source = _load_building_block_library(library_value)
        blocks.extend(library_values)
        library_sources.append(library_source)
    blocks = list(dict.fromkeys(blocks))
    if not blocks:
        raise ValueError("At least one building block or substrate-library entry is required.")

    options = payload.get("config") or {}
    if not isinstance(options, Mapping):
        raise ValueError("config must be a JSON object.")
    options = dict(options)
    ignored = {
        "target", "targets", "target_smiles", "building_blocks", "substrates",
        "building_block_library", "substrate_library", "endpoint_mode", "config",
    }
    aliases = {"top_k": "top_k_routes", "beam_width": "global_beam_width"}
    for key, value in payload.items():
        if key not in ignored:
            options[aliases.get(key, key)] = value
    endpoint_mode = payload.get("endpoint_mode")
    if endpoint_mode is not None:
        endpoint_mode = str(endpoint_mode).strip().casefold()
        aliases_mode = {
            "any": "any_stock", "stock": "any_stock", "library": "any_stock",
            "all": "exact_set", "exact": "exact_set", "required": "exact_set",
        }
        endpoint_mode = aliases_mode.get(endpoint_mode, endpoint_mode)
        if endpoint_mode not in {"any_stock", "exact_set"}:
            raise ValueError("endpoint_mode must be 'any_stock' or 'exact_set'.")
        options["require_all_building_blocks"] = endpoint_mode == "exact_set"
    return targets, blocks, options, {
        "library_sources": library_sources,
        "requested_endpoint_mode": endpoint_mode,
    }


class EnzymeSingleStepEvaluator(BaseTool):
    """Rank externally supplied one-step retrosynthesis candidates."""

    name: str = "EnzymeSingleStepEvaluator"
    description: str = (
        "Evaluate and rank ONE-STEP enzymatic retrosynthesis candidates. Input JSON: "
        '{"target":"product SMILES/name/Cxxxxx","building_blocks":["..."],'
        '"candidates":[{"precursors":["SMILES", "..."],"model_score":0.8,'
        '"feasibility":0.7,"ec":"1.1.1.1","source":"MODEL"}]}. '
        "The candidates field may also be named predictions and accepts raw MooRetro "
        "records with reactants/scores/templates. Ranking uses Morgan-Tanimoto "
        "similarity, MCS target overlap, model/database evidence, feasibility, EC "
        "evidence, and similarity to available building blocks."
    )

    db_path: str = _DEFAULT_DB_PATH
    _database: EnzymeReactionDatabase = PrivateAttr()
    _base_config: PlannerConfig = PrivateAttr()

    def __init__(
        self,
        db_path: Optional[str] = None,
        config: Optional[Mapping[str, Any]] = None,
        database_service: Optional[EnzymeReactionDatabase] = None,
        **kwargs: Any,
    ):
        database = database_service or EnzymeReactionDatabase(db_path or _DEFAULT_DB_PATH)
        base = PlannerConfig.from_mapping(
            config, use_database=True, use_model=False
        )
        super().__init__(db_path=database.db_path, **kwargs)
        self._database = database
        self._base_config = base

    def _candidate(
        self,
        raw: Mapping[str, Any],
        target: MoleculeNode,
        index: int,
    ) -> ReactionCandidate:
        precursor_values = raw.get("precursors")
        if precursor_values is None:
            precursor_values = raw.get("reactants")
        precursor_queries = _string_list(precursor_values)
        if not precursor_queries:
            raise ValueError("candidate has no precursors/reactants")
        precursors = tuple(
            self._database.normalize_molecule(value) for value in precursor_queries
        )
        main, auxiliary, currency = classify_precursors(
            target, precursors, self._base_config.currency_metabolite_ids
        )
        source = str(raw.get("source") or raw.get("provenance") or "MODEL").upper()
        if source not in {"DATABASE", "MODEL", "BOTH"}:
            source = "MODEL"
        ec_value = raw.get("ec_numbers")
        if isinstance(ec_value, (list, tuple)):
            ec_numbers = tuple(str(value).replace("EC ", "") for value in ec_value)
        else:
            ec_numbers = extract_ec_numbers(str(raw.get("ec") or ec_value or ""))
        database_ids = raw.get("database_reaction_ids") or []
        if isinstance(database_ids, str):
            database_ids = [database_ids]
        return ReactionCandidate(
            candidate_id=str(
                raw.get("candidate_id") or f"external:{target.key}:{index + 1}"
            ),
            target=target,
            precursors=precursors,
            source=source,
            main_precursor_keys=main,
            auxiliary_precursor_keys=auxiliary,
            currency_precursor_keys=currency,
            coefficients=tuple("1" for _ in precursors),
            all_products=(target,),
            product_coefficients=("1",),
            database_reaction_ids=tuple(str(value) for value in database_ids),
            reaction_name=raw.get("reaction_name") or raw.get("name"),
            equation=raw.get("stored_equation") or raw.get("equation"),
            ec_numbers=ec_numbers,
            subsystem=raw.get("subsystem"),
            template=raw.get("template") or raw.get("templates"),
            rank=str(raw.get("model_rank") or raw.get("topk") or index + 1),
            database_support=_as_float(raw.get("database_support")),
            model_score=_as_float(
                raw.get("model_score")
                if raw.get("model_score") is not None
                else raw.get("scores", raw.get("score"))
            ),
            structural_similarity=_as_float(
                raw.get("structural_similarity")
                if raw.get("structural_similarity") is not None
                else raw.get("similarity")
            ),
            feasibility=_as_float(raw.get("feasibility")),
            toxicity=_as_float(raw.get("toxicity")),
            ec_consistency=_as_float(raw.get("ec_consistency")),
            metadata={
                "metric_provenance": {
                    "structural_similarity": "RDKIT_MORGAN_TANIMOTO",
                    "approx_structural_overlap": "RDKIT_MCS_TARGET_FRACTION",
                    "building_block_similarity": "RDKIT_MORGAN_TANIMOTO",
                    "external_model_score": "CALLER_SUPPLIED",
                    "external_feasibility": "CALLER_SUPPLIED",
                }
            },
        )

    def _run(self, input_str: str) -> str:
        try:
            payload = _json_object(input_str)
            target = self._database.normalize_molecule(
                _query_from_molecule(
                    payload.get("target")
                    or payload.get("target_smiles")
                    or payload.get("product_smiles")
                )
            )
            block_values = payload.get("building_blocks") or []
            building_blocks = tuple(
                self._database.normalize_molecule(_query_from_molecule(value))
                for value in block_values
            )
            raw_candidates = payload.get("candidates")
            if raw_candidates is None:
                raw_candidates = payload.get("predictions")
            if not isinstance(raw_candidates, list) or not raw_candidates:
                raise ValueError("candidates/predictions must be a non-empty JSON list.")
            candidates = []
            rejected = []
            for index, raw in enumerate(raw_candidates):
                if not isinstance(raw, Mapping):
                    rejected.append({"index": index, "reason": "candidate is not an object"})
                    continue
                try:
                    candidates.append(self._candidate(raw, target, index))
                except ValueError as exc:
                    rejected.append({"index": index, "reason": str(exc)})
            if not candidates:
                raise ValueError("No valid one-step candidates remained after parsing.")
            ranked = score_and_rank_candidates(
                candidates, building_blocks, self._base_config
            )
        except (TypeError, ValueError) as exc:
            return _error("invalid_input", str(exc))
        return json.dumps(
            {
                "ok": True,
                "evaluation": "single_step",
                "target": target.to_dict(),
                "candidate_count": len(ranked),
                "ranked_candidates": [
                    {"rank": index, **candidate.to_dict()}
                    for index, candidate in enumerate(ranked, start=1)
                ],
                "rejected_candidates": rejected,
                "ranking_rule": "descending masked weighted score / ascending cost",
                "config": {
                    "scoring_weights": self._base_config.scoring_weights.to_dict()
                },
            },
            ensure_ascii=False,
        )

    async def _arun(self, input_str: str) -> str:
        raise NotImplementedError("Asynchronous single-step evaluation is not implemented.")


class EnzymeRetrosynthesisSearch(BaseTool):
    """Run model+database AND-OR MCTS for one or more target molecules."""

    name: str = "EnzymeRetrosynthesisSearch"
    description: str = (
        "Plan enzymatic retrosynthesis for one or more targets with a local "
        "MooRetro neural single-step model AND the local reaction database at every "
        "expanded molecule. Input JSON: {\"targets\":[\"SMILES\", ...], "
        "\"building_blocks\":[\"SMILES\", ...], or use "
        "\"building_block_library\":\"/local/file.txt|csv|json|xlsx\", "
        "\"endpoint_mode\":\"any_stock\",\"top_k\":10}. The default "
        "any_stock mode stops only when every non-currency route leaf is present "
        "in the supplied stock/library; exact_set additionally requires every "
        "supplied building block to be used. Every expanded route state first "
        "creates MODEL and DATABASE source branches; each branch owns its reaction "
        "AND nodes and is selected by epsilon-greedy UCB using reaction quality and "
        "stock similarity. Complete routes are then merged and returned as Top-K "
        "routes (or Top-K multi-target route bundles)."
    )

    db_path: str = _DEFAULT_DB_PATH
    _database: EnzymeReactionDatabase = PrivateAttr()
    _model_client: Optional[EnzRetroClient] = PrivateAttr(default=None)
    _base_config: PlannerConfig = PrivateAttr()

    def __init__(
        self,
        db_path: Optional[str] = None,
        config: Optional[Mapping[str, Any]] = None,
        database_service: Optional[EnzymeReactionDatabase] = None,
        model_client: Optional[EnzRetroClient] = None,
        moo_root: Optional[str] = None,
        single_fn_type: str = "biocyc",
        ec_level: int = 3,
        **kwargs: Any,
    ):
        database = database_service or EnzymeReactionDatabase(db_path or _DEFAULT_DB_PATH)
        values = dict(config or {})
        values.setdefault("use_database", True)
        values.setdefault("use_model", True)
        values.setdefault("search_algorithm", "and_or_mcts")
        # Stock/library semantics are the general default. Exact endpoint tests can
        # opt in with endpoint_mode=exact_set or require_all_building_blocks=true.
        values.setdefault("require_all_building_blocks", False)
        base = PlannerConfig.from_mapping(values)
        client = model_client
        if base.use_model and client is None:
            client = EnzRetroClient(
                moo_root=moo_root,
                single_fn_type=single_fn_type,
                ec_level=ec_level,
            )
        super().__init__(db_path=database.db_path, **kwargs)
        self._database = database
        self._model_client = client
        self._base_config = base

    @staticmethod
    def _solution_graph(
        result: Mapping[str, Any], config: PlannerConfig
    ) -> Dict[str, Any]:
        """Build the Top-K chemical graph plus explicit model/database branches."""
        molecules: Dict[str, Dict[str, Any]] = {}
        sources: Dict[str, Dict[str, Any]] = {}
        reactions: Dict[str, Dict[str, Any]] = {}
        edges: Dict[Tuple[str, str, str], Dict[str, str]] = {}
        root = result["target"]["key"]
        stock = {item["key"] for item in result.get("building_blocks", [])}

        def record_edge(source: str, target: str, relation: str) -> None:
            edges[(source, target, relation)] = {
                "source": source,
                "target": target,
                "relation": relation,
            }

        def record_molecule(
            molecule: Mapping[str, Any], route_rank: int, role: str = "intermediate"
        ) -> None:
            key = molecule.get("key")
            if not key:
                return
            node = molecules.setdefault(
                str(key),
                {
                    "node_id": key,
                    "node_type": "OR",
                    "role": role,
                    "molecule": dict(molecule),
                    "solution_routes": [],
                },
            )
            if key == root:
                node["role"] = "target"
            elif key in stock:
                node["role"] = "building_block"
            if route_rank and route_rank not in node["solution_routes"]:
                node["solution_routes"].append(route_rank)

        def record_source(
            molecule_key: str, source: str, route_rank: int = 0
        ) -> str:
            node_id = f"source::{molecule_key}::{source}"
            node = sources.setdefault(
                node_id,
                {
                    "node_id": node_id,
                    "node_type": "SOURCE",
                    "source": source,
                    "molecule_key": molecule_key,
                    "reaction_node_ids": [],
                    "solution_routes": [],
                },
            )
            if route_rank and route_rank not in node["solution_routes"]:
                node["solution_routes"].append(route_rank)
            record_edge(molecule_key, node_id, "SOURCE_PARTITION")
            return node_id

        record_molecule(result["target"], 0, "target")
        enabled_sources = []
        if config.use_model:
            enabled_sources.append("MODEL")
        if config.use_database:
            enabled_sources.append("DATABASE")
        for source in enabled_sources:
            record_source(root, source)

        for route in result.get("routes", []):
            rank = int(route["rank"])
            record_molecule(result["target"], rank, "target")
            for step in route.get("steps", []):
                record_molecule(step["target"], rank)
                precursor_keys = []
                for precursor in step.get("precursors", []):
                    if precursor.get("role") == "currency":
                        continue
                    record_molecule(precursor, rank)
                    precursor_keys.append(precursor["key"])
                candidate_id = step["candidate_id"]
                provenance = str(step.get("provenance") or "UNKNOWN").upper()
                step_sources = (
                    ["MODEL", "DATABASE"]
                    if provenance == "BOTH"
                    else [provenance]
                )
                source_node_ids = [
                    record_source(step["target"]["key"], source, rank)
                    for source in step_sources
                ]
                node = reactions.setdefault(
                    candidate_id,
                    {
                        "node_id": candidate_id,
                        "node_type": "AND",
                        "target_key": step["target"]["key"],
                        "precursor_keys": precursor_keys,
                        "provenance": provenance,
                        "source_branch_node_ids": [],
                        "database_reaction_ids": step.get("database_reaction_ids", []),
                        "ec_numbers": step.get("ec_numbers", []),
                        "step_score": step.get("score"),
                        "step_cost": step.get("cost"),
                        "solution_routes": [],
                    },
                )
                for source_node_id in source_node_ids:
                    if source_node_id not in node["source_branch_node_ids"]:
                        node["source_branch_node_ids"].append(source_node_id)
                    source_node = sources[source_node_id]
                    if candidate_id not in source_node["reaction_node_ids"]:
                        source_node["reaction_node_ids"].append(candidate_id)
                    record_edge(source_node_id, candidate_id, "REACTION_CHOICE")
                for precursor_key in precursor_keys:
                    record_edge(candidate_id, precursor_key, "REQUIRES_ALL")
                if rank not in node["solution_routes"]:
                    node["solution_routes"].append(rank)

        runtime_tree = result.get("mcts_search_tree")
        return {
            "representation": "AND_OR_HYPERGRAPH",
            "source_partitioned": True,
            "root_molecule_key": root,
            "root_source_branches": [
                sources[f"source::{root}::{source}"] for source in enabled_sources
            ],
            "source_branch_order": ["MODEL", "DATABASE"],
            "or_semantics": "choose one source/reaction alternative for each molecule",
            "and_semantics": (
                "solve every non-currency precursor of the selected reaction; "
                "no children[0] shortcut is permitted"
            ),
            "molecule_nodes": list(molecules.values()),
            "source_nodes": list(sources.values()),
            "reaction_nodes": list(reactions.values()),
            "edges": list(edges.values()),
            "molecule_node_count": len(molecules),
            "source_node_count": len(sources),
            "reaction_node_count": len(reactions),
            "runtime_mcts_tree": runtime_tree,
            "scope": "TOP_K_COMPLETE_SOLUTION_SUBGRAPH_PLUS_RUNTIME_MCTS",
        }

    @staticmethod
    def _algorithm_record(algorithm: str, config: PlannerConfig) -> Dict[str, Any]:
        if algorithm == "and_or_mcts":
            node_selection = (
                "MooRetro-style epsilon-greedy UCB1 over explicit MODEL and "
                "DATABASE source branches, with candidate-quality and stock priors"
            )
        elif algorithm == "best_first_beam":
            node_selection = "minimum evaluated cost plus stock-distance heuristic"
        else:
            node_selection = "minimum accumulated evaluated cost"
        return {
            "name": algorithm,
            "graph": "AND_OR_HYPERGRAPH",
            "node_selection": node_selection,
            "candidate_generation": (
                "local MooRetro model and local database queried separately for "
                "every expanded OR state; candidates remain under MODEL/DATABASE "
                "source nodes during MCTS and complete routes are merged only for Top-K"
            ),
            "source_partition": {
                "order": ["MODEL", "DATABASE"],
                "node_type": "SOURCE",
                "applied_at_every_expanded_state": True,
            },
            "mcts": (
                {
                    "selection": "EPSILON_GREEDY_UCB1_OVER_SOURCE_AND_REACTION_NODES",
                    "expansion": (
                        "OR route state -> MODEL/DATABASE source nodes -> reaction "
                        "AND state containing all non-currency precursors"
                    ),
                    "simulation": "bounded source-partitioned evaluator-guided rollout",
                    "backpropagation": "scalar route-quality reward through source nodes",
                    "exploration_constant": config.mcts_exploration_constant,
                    "rollout_depth": config.mcts_rollout_depth,
                    "random_seed": config.mcts_random_seed,
                    "progressive_widening": config.mcts_progressive_widening,
                }
                if algorithm == "and_or_mcts"
                else None
            ),
            "beam_width": config.global_beam_width if algorithm != "and_or_mcts" else None,
        }

    def _finalize_single_result(
        self,
        result: Dict[str, Any],
        algorithm: str,
        config: PlannerConfig,
    ) -> Dict[str, Any]:
        valid_routes = [
            route
            for route in result.get("routes", [])
            if route.get("route_validation", {}).get("valid") is True
        ]
        for rank, route in enumerate(valid_routes, start=1):
            route["rank"] = rank
        result["routes"] = valid_routes
        result["route_count"] = len(valid_routes)
        exact = config.require_all_building_blocks
        result["endpoint_task"] = {
            "type": (
                "target_to_exact_known_substrate_set"
                if exact
                else "target_to_any_supplied_stock_subset"
            ),
            "endpoint_mode": "exact_set" if exact else "any_stock",
            "rooted": True,
            "connected": True,
            "acyclic": True,
            "all_terminal_molecules_must_be_in_stock": True,
            "terminal_set_must_equal_building_blocks": exact,
        }
        result["candidate_generation"] = {
            "database_enabled": config.use_database,
            "model_enabled": config.use_model,
            "model_backend": (
                self._model_client.cache_info().get("backend_mode")
                if self._model_client is not None
                else None
            ),
            "source_branches": [
                source
                for source, enabled in (
                    ("MODEL", config.use_model),
                    ("DATABASE", config.use_database),
                )
                if enabled
            ],
            "fusion": (
                "SOURCE_PARTITIONED_DURING_MCTS; COMPLETE_ROUTES_MERGED_FOR_TOP_K"
                if algorithm == "and_or_mcts"
                else "UNION_WITH_EXACT_PRECURSOR_AND_EC_CONSENSUS"
            ),
            "performed_at_each_expanded_molecule": True,
        }
        result["planner"] = self.name
        result["search_algorithm"] = self._algorithm_record(algorithm, config)
        result["and_or_tree"] = self._solution_graph(result, config)
        return result

    @staticmethod
    def _route_score(route: Mapping[str, Any]) -> float:
        value = route.get("diversity_adjusted_score", route.get("route_score", 0.0))
        try:
            return max(1.0e-12, float(value))
        except (TypeError, ValueError):
            return 1.0e-12

    @classmethod
    def _top_route_bundles(
        cls,
        target_results: Sequence[Mapping[str, Any]],
        top_k: int,
    ) -> List[Dict[str, Any]]:
        route_lists = [
            sorted(
                list(result.get("routes") or ()),
                key=lambda route: (-cls._route_score(route), int(route.get("rank", 0))),
            )
            for result in target_results
        ]
        if not route_lists or any(not routes for routes in route_lists):
            return []

        def priority(indices: Tuple[int, ...]) -> Tuple[float, int, Tuple[int, ...]]:
            selected = [route_lists[i][index] for i, index in enumerate(indices)]
            negative_log_score = -sum(math.log(cls._route_score(route)) for route in selected)
            step_count = sum(int(route.get("step_count", 0)) for route in selected)
            return negative_log_score, step_count, indices

        initial = tuple(0 for _ in route_lists)
        heap = [priority(initial)]
        visited = {initial}
        bundles = []
        while heap and len(bundles) < top_k:
            _, total_steps, indices = heapq.heappop(heap)
            selected = [route_lists[i][index] for i, index in enumerate(indices)]
            combined_score = math.exp(
                sum(math.log(cls._route_score(route)) for route in selected)
                / len(selected)
            )
            bundles.append(
                {
                    "rank": len(bundles) + 1,
                    "complete": True,
                    "combined_route_score": round(combined_score, 6),
                    "total_step_count": total_steps,
                    "target_routes": [
                        {
                            "target": dict(target_results[i]["target"]),
                            "route_rank": int(route.get("rank", 0)),
                            "route": route,
                        }
                        for i, route in enumerate(selected)
                    ],
                    "ranking_rule": (
                        "geometric mean of per-target diversity-adjusted route scores, "
                        "then total step count"
                    ),
                }
            )
            for dimension in range(len(indices)):
                neighbor = list(indices)
                neighbor[dimension] += 1
                neighbor_tuple = tuple(neighbor)
                if (
                    neighbor[dimension] < len(route_lists[dimension])
                    and neighbor_tuple not in visited
                ):
                    visited.add(neighbor_tuple)
                    heapq.heappush(heap, priority(neighbor_tuple))
        return bundles

    def _run(self, input_str: str) -> str:
        try:
            targets, building_blocks, overrides, input_metadata = _parse_search_input(input_str)
            algorithm = str(
                overrides.pop("search_algorithm", self._base_config.search_algorithm)
            )
            if algorithm not in {"best_first_beam", "uniform_cost", "and_or_mcts"}:
                raise ValueError(
                    "search_algorithm must be best_first_beam, uniform_cost, or and_or_mcts."
                )
            merged = self._base_config.to_dict()
            merged.update(overrides)
            merged["search_algorithm"] = algorithm
            if algorithm == "uniform_cost":
                merged["heuristic_weight"] = 0.0
            config = PlannerConfig.from_mapping(merged)
        except (TypeError, ValueError, OSError, json.JSONDecodeError) as exc:
            return _error("invalid_input", str(exc))

        target_results = []
        try:
            for target in targets:
                result = HybridPlannerEngine(
                    database=self._database,
                    model_client=self._model_client,
                    config=config,
                ).plan(target, building_blocks)
                if not result.get("ok"):
                    return json.dumps(result, ensure_ascii=False)
                target_results.append(
                    self._finalize_single_result(result, algorithm, config)
                )
        except Exception as exc:
            return _error("search_failed", str(exc))

        if len(target_results) == 1:
            result = target_results[0]
            result["input"] = {
                "target_count": 1,
                "building_block_count": len(result.get("building_blocks", [])),
                **input_metadata,
            }
            return json.dumps(result, ensure_ascii=False)

        route_bundles = self._top_route_bundles(
            target_results, config.top_k_routes
        )
        output = {
            "ok": True,
            "planner": self.name,
            "mode": "multi-target",
            "targets": [result["target"] for result in target_results],
            "target_count": len(target_results),
            "building_blocks": target_results[0].get("building_blocks", []),
            "building_block_count": len(target_results[0].get("building_blocks", [])),
            "endpoint_task": target_results[0]["endpoint_task"],
            "candidate_generation": target_results[0]["candidate_generation"],
            "search_algorithm": self._algorithm_record(algorithm, config),
            "all_targets_solved": all(result["route_count"] > 0 for result in target_results),
            "target_results": target_results,
            "route_count": len(route_bundles),
            "routes": route_bundles,
            "route_bundles": route_bundles,
            "and_or_forest": [result["and_or_tree"] for result in target_results],
            "stats": {
                "per_target": [result.get("stats", {}) for result in target_results],
                "model_calls": sum(
                    int(result.get("stats", {}).get("model_calls", 0))
                    for result in target_results
                ),
                "database_candidates": sum(
                    int(result.get("stats", {}).get("database_candidates", 0))
                    for result in target_results
                ),
                "model_candidates": sum(
                    int(result.get("stats", {}).get("model_candidates", 0))
                    for result in target_results
                ),
                "returned_route_bundle_count": len(route_bundles),
            },
            "config": config.to_dict(),
            "input": input_metadata,
        }
        return json.dumps(output, ensure_ascii=False)

    async def _arun(self, input_str: str) -> str:
        raise NotImplementedError("Asynchronous retrosynthesis search is not implemented.")


class EnzymeMultiStepEvaluator(BaseTool):
    """Re-score complete routes and return a diversity-aware Top-K list."""

    name: str = "EnzymeMultiStepEvaluator"
    description: str = (
        "Evaluate and Top-K rank COMPLETE multi-step retrosynthesis routes. Input "
        'JSON: {"routes":[planner route objects],"top_k":5,'
        '"diversity_penalty":0.1}. It combines geometric-mean step quality, route '
        "completion, route length, approximate carbon efficiency, organic byproduct "
        "burden and cofactor burden, then applies reaction-set diversity reranking. "
        "The routes returned by EnzymeRetrosynthesisSearch can be passed directly."
    )

    _default_weights: Dict[str, float] = PrivateAttr()

    def __init__(self, **kwargs: Any):
        super().__init__(**kwargs)
        self._default_weights = {
            "reaction_quality": 0.55,
            "carbon_efficiency": 0.15,
            "inverse_byproduct_burden": 0.10,
            "inverse_cofactor_burden": 0.05,
            "route_length": 0.10,
            "completion": 0.05,
        }

    @staticmethod
    def _step_quality(route: Mapping[str, Any]) -> float:
        values = []
        for step in route.get("steps", []):
            score = _as_float(step.get("score"))
            if score is None:
                metrics = step.get("raw_metrics") or {}
                score = _as_float(
                    metrics.get("model_score")
                    if metrics.get("model_score") is not None
                    else metrics.get("database_support")
                )
            if score is not None:
                values.append(max(1.0e-8, _clamp(score)))
        if values:
            return math.exp(sum(math.log(value) for value in values) / len(values))
        fallback = _as_float(route.get("route_score"))
        return _clamp(fallback) if fallback is not None else 0.0

    def _evaluate(
        self, route: Mapping[str, Any], weights: Mapping[str, float]
    ) -> Dict[str, Any]:
        step_count = int(route.get("step_count", len(route.get("steps", []))))
        complete = bool(
            route.get("complete", True)
            and route.get("all_noncurrency_leaves_solved", True)
        )
        components: Dict[str, Optional[float]] = {
            "reaction_quality": self._step_quality(route),
            "carbon_efficiency": _as_float(
                route.get("target_carbon_efficiency")
            ),
            "inverse_byproduct_burden": (
                None
                if route.get("major_organic_byproduct_penalty") is None
                else 1.0
                - _clamp(float(route["major_organic_byproduct_penalty"]))
            ),
            "inverse_cofactor_burden": (
                None
                if route.get("currency_cofactor_count") is None
                else 1.0 / (1.0 + max(0.0, float(route["currency_cofactor_count"])))
            ),
            "route_length": 1.0 / (1.0 + 0.15 * max(0, step_count - 1)),
            "completion": 1.0 if complete else 0.0,
        }
        active = {
            name: _clamp(value)
            for name, value in components.items()
            if value is not None and weights.get(name, 0.0) > 0.0
        }
        denominator = sum(weights[name] for name in active)
        score = (
            sum(active[name] * weights[name] for name in active) / denominator
            if denominator > 0
            else 0.0
        )
        return {
            "route": dict(route),
            "base_score": _clamp(score),
            "components": active,
            "signature": _route_signature(route),
        }

    def _run(self, input_str: str) -> str:
        try:
            payload = _json_object(input_str)
            routes = payload.get("routes")
            if not isinstance(routes, list) or not routes:
                raise ValueError("routes must be a non-empty JSON list.")
            top_k = int(payload.get("top_k", len(routes)))
            if top_k < 1:
                raise ValueError("top_k must be at least 1.")
            diversity_penalty = float(payload.get("diversity_penalty", 0.10))
            if not 0.0 <= diversity_penalty <= 1.0:
                raise ValueError("diversity_penalty must be between 0 and 1.")
            weights = dict(self._default_weights)
            custom_weights = payload.get("weights") or {}
            if not isinstance(custom_weights, Mapping):
                raise ValueError("weights must be a JSON object.")
            unknown = sorted(set(custom_weights) - set(weights))
            if unknown:
                raise ValueError("Unknown route weight(s): " + ", ".join(unknown))
            weights.update({key: float(value) for key, value in custom_weights.items()})
            if any(value < 0 for value in weights.values()) or sum(weights.values()) <= 0:
                raise ValueError("Route weights must be non-negative with a positive sum.")
            remaining = [
                self._evaluate(route, weights)
                for route in routes
                if isinstance(route, Mapping)
            ]
            if not remaining:
                raise ValueError("No valid route objects were supplied.")
        except (TypeError, ValueError) as exc:
            return _error("invalid_input", str(exc))

        selected = []
        while remaining and len(selected) < top_k:
            best_index = 0
            best_key = (-float("inf"), -float("inf"))
            for index, value in enumerate(remaining):
                similarity = max(
                    (
                        _jaccard(value["signature"], chosen["signature"])
                        for chosen in selected
                    ),
                    default=0.0,
                )
                adjusted = value["base_score"] - diversity_penalty * similarity
                key = (adjusted, value["base_score"])
                if key > best_key:
                    best_index = index
                    best_key = key
            chosen = remaining.pop(best_index)
            chosen["diversity_adjusted_score"] = best_key[0]
            selected.append(chosen)

        ranked_routes = []
        for rank, value in enumerate(selected, start=1):
            route = value["route"]
            route["evaluation_rank"] = rank
            route["evaluation_score"] = round(value["base_score"], 6)
            route["evaluation_diversity_adjusted_score"] = round(
                value["diversity_adjusted_score"], 6
            )
            route["evaluation_components"] = {
                key: round(component, 6)
                for key, component in value["components"].items()
            }
            ranked_routes.append(route)
        return json.dumps(
            {
                "ok": True,
                "evaluation": "multi_step",
                "route_count": len(ranked_routes),
                "top_k_routes": ranked_routes,
                "weights": weights,
                "diversity_penalty": diversity_penalty,
                "aggregation": (
                    "masked weighted average; reaction quality is geometric mean "
                    "of step scores"
                ),
                "metric_provenance": {
                    "reaction_quality": "GEOMETRIC_MEAN_OF_STEP_SCORES",
                    "carbon_efficiency": "PLANNER_STOICHIOMETRIC_CARBON_APPROXIMATION",
                    "route_length": "MONOTONIC_STEP_COUNT_PENALTY",
                    "diversity": "JACCARD_DISTANCE_OVER_REACTION_IDENTITIES",
                },
            },
            ensure_ascii=False,
        )

    async def _arun(self, input_str: str) -> str:
        raise NotImplementedError("Asynchronous route evaluation is not implemented.")
